# import re
# from docx import Document
# from docx.oxml.ns import qn


# class DOCXExtractor:

#     @staticmethod
#     def extract_text(docx_path):
#         try:
#             doc = Document(docx_path)
#             text_parts = []
#             seen = set()

#             # --- Headers & Footers ---
#             for section in doc.sections:
#                 for container in [section.header, section.footer]:
#                     if container is None:
#                         continue
#                     for para in container.paragraphs:
#                         t = para.text.strip()
#                         if t:
#                             key = DOCXExtractor._normalize(t)
#                             if key not in seen:
#                                 seen.add(key)
#                                 text_parts.append(t)

#             # --- Body: Paragraphs + Tables ---
#             for element in doc.element.body:
#                 tag = element.tag

#                 if tag.endswith("}p"):
#                     para_text = DOCXExtractor._para_text(element)
#                     if not para_text.strip():
#                         continue
#                     key = DOCXExtractor._normalize(para_text)
#                     if key in seen:
#                         continue
#                     seen.add(key)
#                     text_parts.append(para_text.strip())

#                 elif tag.endswith("}tbl"):
#                     for row in element.iter(qn("w:tr")):
#                         row_text = DOCXExtractor._process_row(row)
#                         if not row_text:
#                             continue
#                         key = DOCXExtractor._normalize(row_text)
#                         if key in seen:
#                             continue
#                         seen.add(key)
#                         text_parts.append(row_text)

#             # --- Text Boxes / Shapes / Drawings ---
#             shape_text = DOCXExtractor._extract_all_shape_text(doc)
#             if shape_text:
#                 for line in shape_text.split("\n"):
#                     line = line.strip()
#                     if line:
#                         key = DOCXExtractor._normalize(line)
#                         if key not in seen:
#                             seen.add(key)
#                             text_parts.append(line)

#             # --- BRUTE FORCE FALLBACK ---
#             # If we still have very little text, dump ALL w:t nodes
#             if len("\n".join(text_parts)) < 100:
#                 brute_text = DOCXExtractor._brute_force_text(doc)
#                 if brute_text:
#                     for line in brute_text.split("\n"):
#                         line = line.strip()
#                         if line:
#                             key = DOCXExtractor._normalize(line)
#                             if key not in seen:
#                                 seen.add(key)
#                                 text_parts.append(line)

#             return "\n".join(text_parts).strip()

#         except Exception as e:
#             print(f"DOCX Extraction Error: {e}")
#             return ""

#     @staticmethod
#     def _extract_all_shape_text(doc):
#         """Extract text from all shapes, text boxes, drawings"""
#         shape_texts = []
#         for para in doc.element.iter(qn("w:p")):
#             parent = para.getparent()
#             in_shape = False
#             while parent is not None:
#                 tag = parent.tag
#                 if tag.endswith("}drawing") or tag.endswith("}pict") or tag.endswith("}txbxContent"):
#                     in_shape = True
#                     break
#                 parent = parent.getparent()

#             if in_shape:
#                 text = DOCXExtractor._para_text(para)
#                 if text.strip():
#                     shape_texts.append(text.strip())
#         return "\n".join(shape_texts) if shape_texts else ""

#     @staticmethod
#     def _brute_force_text(doc):
#         """Last resort: grab every w:t node in the entire document"""
#         texts = []
#         for t in doc.element.iter(qn("w:t")):
#             if t.text and t.text.strip():
#                 texts.append(t.text.strip())
#         # Reconstruct with simple joining
#         raw = " ".join(texts)
#         # Try to split by obvious paragraph markers
#         return raw.replace("  ", "\n").replace(". ", ".\n")

#     @staticmethod
#     def _para_text(element):
#         return "".join(node.text for node in element.iter(qn("w:t")) if node.text)

#     @staticmethod
#     def _process_row(row):
#         seen_cells = set()
#         row_cells = []
#         for cell in row.iter(qn("w:tc")):
#             cell_text = DOCXExtractor._cell_text(cell)
#             if not cell_text.strip():
#                 continue
#             key = DOCXExtractor._normalize(cell_text)
#             if key in seen_cells:
#                 continue
#             seen_cells.add(key)
#             row_cells.append(cell_text.strip())
#         if not row_cells:
#             return ""
#         return " | ".join(row_cells)

#     @staticmethod
#     def _cell_text(cell):
#         paras = []
#         for para in cell.iter(qn("w:p")):
#             t = "".join(node.text for node in para.iter(qn("w:t")) if node.text)
#             if t.strip():
#                 paras.append(t.strip())
#         return "\n".join(paras)

#     @staticmethod
#     def _normalize(text):
#         return re.sub(r"\s+", " ", text.strip().lower())

import re
from docx import Document
from docx.oxml.ns import qn


class DOCXExtractor:

    @staticmethod
    def extract_text(docx_path):
        try:
            doc = Document(docx_path)
            text_parts = []
            seen = set()

            # --- Headers & Footers (all variants) ---
            for section in doc.sections:
                for container in [
                    section.header,
                    section.footer,
                    section.even_page_header,
                    section.even_page_footer,
                    section.first_page_header,
                    section.first_page_footer,
                ]:
                    if container is None:
                        continue

                    # Normal paragraphs in header/footer
                    for para in container.paragraphs:
                        t = para.text.strip()
                        if t:
                            key = DOCXExtractor._normalize(t)
                            if key not in seen:
                                seen.add(key)
                                text_parts.append(t)

                    # =========================================
                    # FIX: Text boxes inside headers/footers
                    # Many DOCX files put contact info in
                    # text boxes inside the header, not in
                    # regular paragraphs — these were missed
                    # =========================================
                    for elem in container._element.iter():
                        tag = elem.tag
                        if tag.endswith("}txbxContent"):
                            for p in elem.iter(qn("w:p")):
                                t = DOCXExtractor._para_text(p).strip()
                                if t:
                                    key = DOCXExtractor._normalize(t)
                                    if key not in seen:
                                        seen.add(key)
                                        text_parts.append(t)

            # --- Body: Paragraphs + Tables ---
            for element in doc.element.body:
                tag = element.tag

                if tag.endswith("}p"):
                    para_text = DOCXExtractor._para_text(element)
                    if not para_text.strip():
                        continue
                    key = DOCXExtractor._normalize(para_text)
                    if key in seen:
                        continue
                    seen.add(key)
                    text_parts.append(para_text.strip())

                elif tag.endswith("}tbl"):
                    for row in element.iter(qn("w:tr")):
                        row_text = DOCXExtractor._process_row(row)
                        if not row_text:
                            continue
                        key = DOCXExtractor._normalize(row_text)
                        if key in seen:
                            continue
                        seen.add(key)
                        text_parts.append(row_text)

            # --- Text Boxes / Shapes / Drawings ---
            shape_text = DOCXExtractor._extract_all_shape_text(doc)
            if shape_text:
                for line in shape_text.split("\n"):
                    line = line.strip()
                    if line:
                        key = DOCXExtractor._normalize(line)
                        if key not in seen:
                            seen.add(key)
                            text_parts.append(line)

            # --- BRUTE FORCE FALLBACK ---
            # If we still have very little text, dump ALL w:t nodes
            if len("\n".join(text_parts)) < 100:
                brute_text = DOCXExtractor._brute_force_text(doc)
                if brute_text:
                    for line in brute_text.split("\n"):
                        line = line.strip()
                        if line:
                            key = DOCXExtractor._normalize(line)
                            if key not in seen:
                                seen.add(key)
                                text_parts.append(line)

            return "\n".join(text_parts).strip()

        except Exception as e:
            print(f"DOCX Extraction Error: {e}")
            return ""

    @staticmethod
    def _extract_all_shape_text(doc):
        """Extract text from all shapes, text boxes, drawings"""
        shape_texts = []
        for para in doc.element.iter(qn("w:p")):
            parent = para.getparent()
            in_shape = False
            while parent is not None:
                tag = parent.tag
                if (
                    tag.endswith("}drawing")
                    or tag.endswith("}pict")
                    or tag.endswith("}txbxContent")
                ):
                    in_shape = True
                    break
                parent = parent.getparent()

            if in_shape:
                text = DOCXExtractor._para_text(para)
                if text.strip():
                    shape_texts.append(text.strip())

        return "\n".join(shape_texts) if shape_texts else ""

    @staticmethod
    def _brute_force_text(doc):
        """Last resort: grab every w:t node in the entire document"""
        texts = []
        for t in doc.element.iter(qn("w:t")):
            if t.text and t.text.strip():
                texts.append(t.text.strip())
        raw = " ".join(texts)
        return raw.replace("  ", "\n").replace(". ", ".\n")

    @staticmethod
    def _para_text(element):
        return "".join(
            node.text for node in element.iter(qn("w:t")) if node.text
        )

    @staticmethod
    def _process_row(row):
        seen_cells = set()
        row_cells = []
        for cell in row.iter(qn("w:tc")):
            cell_text = DOCXExtractor._cell_text(cell)
            if not cell_text.strip():
                continue
            key = DOCXExtractor._normalize(cell_text)
            if key in seen_cells:
                continue
            seen_cells.add(key)
            row_cells.append(cell_text.strip())
        if not row_cells:
            return ""
        return " | ".join(row_cells)

    @staticmethod
    def _cell_text(cell):
        paras = []
        for para in cell.iter(qn("w:p")):
            t = "".join(
                node.text for node in para.iter(qn("w:t")) if node.text
            )
            if t.strip():
                paras.append(t.strip())
        return "\n".join(paras)

    @staticmethod
    def _normalize(text):
        return re.sub(r"\s+", " ", text.strip().lower())